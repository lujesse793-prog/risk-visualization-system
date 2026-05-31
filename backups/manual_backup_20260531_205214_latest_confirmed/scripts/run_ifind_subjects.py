import argparse
import html
import json
import os
import re
import time
from collections import OrderedDict
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook


WORKSPACE = Path(__file__).resolve().parents[1]
DATA_DIR = WORKSPACE / "ifind_runs"
EXCEL_DIR = WORKSPACE / "data" / "raw"
MCP_URL = "https://api-mcp.51ifind.com:8643/ds-mcp-servers/hexin-ifind-ds-news-mcp"
MIN_REQUEST_INTERVAL_SECONDS = 0.55

NEGATIVE_RULES = [
    ("失信被执行", 100, ["失信", "失信被执行", "失信名单"]),
    ("债务违约", 100, ["债务违约", "债券违约", "本息违约", "实质性违约"]),
    ("兑付困难", 100, ["兑付困难", "兑付危机", "兑付风险", "到期未兑付"]),
    ("破产重整", 98, ["破产重整", "破产清算", "破产申请", "宣告破产"]),
    ("新增被执行", 92, ["被执行", "执行信息", "执行标的", "新增执行"]),
    ("限制高消费", 85, ["限制高消费", "限高", "限制消费"]),
    ("资产冻结", 82, ["资产冻结", "账户冻结", "股权冻结", "资金冻结"]),
    ("查封", 80, ["查封", "财产查封"]),
    ("评级下调", 78, ["评级下调", "信用下调", "展望负面", "降级"]),
    ("担保代偿", 78, ["担保代偿", "担保履约", "代偿义务"]),
    ("重大诉讼", 75, ["重大诉讼", "诉讼进展", "诉讼公告"]),
    ("仲裁", 73, ["仲裁", "商事仲裁"]),
    ("展期", 65, ["展期", "延期兑付", "延期偿还"]),
    ("评级关注", 60, ["评级关注", "列入观察", "评级展望调整"]),
    ("重大行政处罚", 58, ["重大行政处罚", "监管处罚", "行政处罚"]),
    ("工程款争议", 55, ["工程款", "支付争议", "合同纠纷"]),
    ("经营异常", 45, ["经营异常", "停产", "停工", "重大亏损"]),
    ("区域融资压力", 40, ["融资压力", "再融资困难", "融资环境"]),
]


def post_json(payload, token):
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = Request(
        MCP_URL,
        data=body,
        method="POST",
        headers={
            "Authorization": token,
            "Accept": "application/json, text/event-stream",
            "Content-Type": "application/json",
        },
    )
    with urlopen(req, timeout=90) as response:
        raw = response.read().decode("utf-8", errors="replace")
    return json.loads(raw)


class RateLimiter:
    def __init__(self, min_interval):
        self.min_interval = min_interval
        self.last_request_at = 0.0

    def wait(self):
        now = time.monotonic()
        wait_seconds = self.min_interval - (now - self.last_request_at)
        if wait_seconds > 0:
            time.sleep(wait_seconds)
        self.last_request_at = time.monotonic()


def call_mcp(payload, token, limiter, attempts=4):
    last_error = None
    for attempt in range(1, attempts + 1):
        try:
            limiter.wait()
            return post_json(payload, token)
        except (HTTPError, URLError, TimeoutError, json.JSONDecodeError) as exc:
            last_error = exc
            if attempt == attempts:
                break
            time.sleep(min(8.0, 1.5 * attempt))
    raise last_error


def load_entities(limit=None):
    excel_files = sorted(EXCEL_DIR.glob("*.xlsx"))
    if not excel_files:
        raise FileNotFoundError(f"No xlsx files found in {EXCEL_DIR}")

    workbook = load_workbook(excel_files[0], data_only=True, read_only=True)
    worksheet = workbook.active
    entities = OrderedDict()
    project_count = 0
    borrower_records = 0
    guarantor_records = 0

    for project, borrower, guarantor in worksheet.iter_rows(min_row=2, max_col=3, values_only=True):
        if not project and not borrower and not guarantor:
            continue
        project_count += 1
        for role, name in (("借款人", borrower), ("保证人", guarantor)):
            if not name:
                continue
            name = str(name).strip()
            if name in ("#N/A", "N/A"):
                continue
            if role == "借款人":
                borrower_records += 1
            else:
                guarantor_records += 1
            item = entities.setdefault(name, {"name": name, "roles": set(), "projects": []})
            item["roles"].add(role)
            if project:
                item["projects"].append(str(project).strip())
        if limit and len(entities) >= limit:
            break

    result = []
    for item in list(entities.values())[:limit]:
        result.append(
            {
                "name": item["name"],
                "roles": sorted(item["roles"]),
                "projects": item["projects"],
            }
        )
    return {
        "excel_path": str(excel_files[0]),
        "project_count": project_count,
        "borrower_records": borrower_records,
        "guarantor_records": guarantor_records,
        "entities": result,
    }


def build_query(entity):
    return (
        f"{entity['name']} 近三个月 重大负面舆情 风险事件 "
        "失信 被执行 限制高消费 债务违约 展期 兑付困难 破产重整 重大诉讼 仲裁 "
        "资产冻结 查封 行政处罚 监管处罚 评级下调 评级关注 担保履约 融资压力"
    )


def extract_text(result):
    try:
        content = result["result"]["content"]
        texts = []
        for item in content:
            if item.get("type") == "text":
                texts.append(item.get("text", ""))
        return "\n".join(texts)
    except Exception:
        return json.dumps(result, ensure_ascii=False)


def normalize_name(text):
    return re.sub(r"[\s（）()【】\\[\\]，,。.;；:：\"'“”‘’]", "", text or "")


def parse_ifind_items(text):
    try:
        outer = json.loads(text)
        data = outer.get("data", {})
        raw_items = data.get("data")
        if isinstance(raw_items, str):
            return json.loads(raw_items)
        if isinstance(raw_items, list):
            return raw_items
    except Exception:
        return []
    return []


def relevant_text_for_entity(text, entity_name):
    items = parse_ifind_items(text)
    if not items:
        return text, False

    normalized_entity = normalize_name(entity_name)
    relevant = []
    for item in items:
        blob = " ".join(str(value or "") for value in item.values())
        if normalized_entity and normalized_entity in normalize_name(blob):
            relevant.append(blob)

    if not relevant:
        return "", False
    return "\n".join(relevant), True


def classify_text(text):
    hits = []
    lowered = text.lower()
    for event_type, base, keywords in NEGATIVE_RULES:
        matched = [kw for kw in keywords if kw.lower() in lowered]
        if matched:
            hits.append({"eventType": event_type, "score": base, "keywords": matched})

    if not hits:
        return {
            "level": "无",
            "highRelevance": False,
            "score": 0,
            "eventType": "未发现重大负面舆情",
            "keywords": [],
        }

    best = max(hits, key=lambda item: item["score"])
    score = best["score"]
    if score >= 82:
        level = "严重"
    elif score >= 60:
        level = "较重"
    else:
        level = "关注"
    return {
        "level": level,
        "highRelevance": score >= 60,
        "score": score,
        "eventType": best["eventType"],
        "keywords": sorted({kw for hit in hits for kw in hit["keywords"]}),
    }


def classify_entity_text(text, entity_name):
    relevant_text, exact_match = relevant_text_for_entity(text, entity_name)
    if not exact_match:
        return {
            "level": "无",
            "highRelevance": False,
            "score": 0,
            "eventType": "未发现主体精确匹配重大负面舆情",
            "keywords": [],
            "exactMatch": False,
        }
    classification = classify_text(relevant_text)
    classification["exactMatch"] = True
    return classification


def compact_preview(text, max_len=600):
    text = re.sub(r"\s+", " ", text or "").strip()
    if len(text) <= max_len:
        return text
    return text[:max_len].rstrip() + "..."


def render_report(run):
    rows = run["rows"]
    counts = {level: sum(1 for row in rows if row["classification"]["level"] == level) for level in ["严重", "较重", "关注", "无"]}
    generated = html.escape(run["generated_at"])
    row_html = []
    for row in rows:
        c = row["classification"]
        raw_rel = Path(row["raw_path"]).name
        projects = "<br>".join(html.escape(p) for p in row["entity"]["projects"][:3])
        if len(row["entity"]["projects"]) > 3:
            projects += f"<br>等 {len(row['entity']['projects'])} 个项目"
        row_html.append(
            "<tr>"
            f"<td>{row['index']}</td>"
            f"<td class='entity'>{html.escape(row['entity']['name'])}</td>"
            f"<td>{html.escape(','.join(row['entity']['roles']))}</td>"
            f"<td>{projects}</td>"
            f"<td><span class='badge {c['level']}'>{html.escape(c['level'])}</span></td>"
            f"<td>{html.escape(c['eventType'])}</td>"
            f"<td>{html.escape(', '.join(c['keywords']))}</td>"
            f"<td>{html.escape(row['preview'])}</td>"
            f"<td><a href='./raw/{raw_rel}'>raw</a></td>"
            "</tr>"
        )

    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>iFinD 风控全量试跑报告</title>
  <style>
    body {{ margin: 0; font-family: "Microsoft YaHei UI", Arial, sans-serif; background: #f6f8fb; color: #182230; }}
    header {{ padding: 28px 34px; background: #132238; color: white; }}
    h1 {{ margin: 0 0 10px; font-size: 24px; }}
    .meta {{ color: #cbd5e1; font-size: 13px; line-height: 1.8; }}
    main {{ padding: 24px 34px 40px; }}
    .stats {{ display: grid; grid-template-columns: repeat(5, minmax(120px, 1fr)); gap: 12px; margin-bottom: 18px; }}
    .stat {{ background: white; border: 1px solid #e2e8f0; border-radius: 8px; padding: 14px; }}
    .stat strong {{ display: block; font-size: 24px; }}
    table {{ width: 100%; border-collapse: collapse; background: white; border: 1px solid #e2e8f0; font-size: 12px; }}
    th, td {{ border-bottom: 1px solid #e2e8f0; padding: 9px 10px; vertical-align: top; text-align: left; }}
    th {{ position: sticky; top: 0; background: #eef3f8; z-index: 1; }}
    .entity {{ font-weight: 700; min-width: 220px; }}
    .badge {{ display: inline-flex; min-width: 44px; justify-content: center; border-radius: 999px; padding: 3px 8px; font-weight: 700; }}
    .严重 {{ color: #b91c1c; background: #fee2e2; }}
    .较重 {{ color: #c2410c; background: #ffedd5; }}
    .关注 {{ color: #a16207; background: #fef3c7; }}
    .无 {{ color: #047857; background: #d1fae5; }}
    a {{ color: #155eef; }}
  </style>
</head>
<body>
  <header>
    <h1>iFinD 风控全量试跑报告</h1>
    <div class="meta">
      生成时间：{generated}<br>
      数据底稿：{html.escape(run['excel_path'])}<br>
      MCP：hexin-ifind-ds-news-mcp / search_news；限速：每次请求间隔不少于 {MIN_REQUEST_INTERVAL_SECONDS} 秒，不超过 2 次/秒。
    </div>
  </header>
  <main>
    <section class="stats">
      <div class="stat"><span>去重主体</span><strong>{len(rows)}</strong></div>
      <div class="stat"><span>严重</span><strong>{counts['严重']}</strong></div>
      <div class="stat"><span>较重</span><strong>{counts['较重']}</strong></div>
      <div class="stat"><span>关注</span><strong>{counts['关注']}</strong></div>
      <div class="stat"><span>无</span><strong>{counts['无']}</strong></div>
    </section>
    <table>
      <thead>
        <tr><th>#</th><th>主体</th><th>角色</th><th>关联项目</th><th>等级</th><th>事件类型</th><th>命中词</th><th>iFinD 返回预览</th><th>原始</th></tr>
      </thead>
      <tbody>
        {''.join(row_html)}
      </tbody>
    </table>
  </main>
</body>
</html>"""


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int)
    parser.add_argument("--size", type=int, default=3)
    args = parser.parse_args()

    token = os.environ.get("HEXIN_IFIND_AUTHORIZATION")
    if not token:
        raise RuntimeError("HEXIN_IFIND_AUTHORIZATION is not set")

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    raw_dir = DATA_DIR / "raw"
    raw_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_path = DATA_DIR / f"ifind_full_run_{stamp}.json"
    report_path = DATA_DIR / f"ifind_full_run_{stamp}.html"

    limiter = RateLimiter(MIN_REQUEST_INTERVAL_SECONDS)
    init_payload = {
        "jsonrpc": "2.0",
        "id": 1,
        "method": "initialize",
        "params": {
            "protocolVersion": "2024-11-05",
            "capabilities": {},
            "clientInfo": {"name": "risk-visualization-ifind-runner", "version": "1.0.0"},
        },
    }
    call_mcp(init_payload, token, limiter)

    loaded = load_entities(args.limit)
    start = (date.today() - timedelta(days=90)).isoformat()
    end = date.today().isoformat()
    rows = []
    total = len(loaded["entities"])
    print(f"Loaded {total} unique entities from {loaded['excel_path']}", flush=True)

    for index, entity in enumerate(loaded["entities"], start=1):
        query = build_query(entity)
        payload = {
            "jsonrpc": "2.0",
            "id": index + 1,
            "method": "tools/call",
            "params": {
                "name": "search_news",
                "arguments": {
                    "query": query,
                    "time_start": start,
                    "time_end": end,
                    "size": args.size,
                },
            },
        }
        print(f"[{index}/{total}] {entity['name']}", flush=True)
        error = None
        try:
            result = call_mcp(payload, token, limiter)
        except (HTTPError, URLError, TimeoutError, json.JSONDecodeError) as exc:
            result = {"error": str(exc)}
            error = str(exc)

        raw_path = raw_dir / f"{index:03d}.json"
        raw_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        text = extract_text(result)
        classification = classify_entity_text(text, entity["name"])
        rows.append(
            {
                "index": index,
                "entity": entity,
                "query": query,
                "error": error,
                "raw_path": str(raw_path),
                "preview": compact_preview(text),
                "classification": classification,
            }
        )

    run = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "excel_path": loaded["excel_path"],
        "project_count": loaded["project_count"],
        "borrower_records": loaded["borrower_records"],
        "guarantor_records": loaded["guarantor_records"],
        "rows": rows,
    }
    run_path.write_text(json.dumps(run, ensure_ascii=False, indent=2), encoding="utf-8")
    report_path.write_text(render_report(run), encoding="utf-8")
    print(f"JSON: {run_path}", flush=True)
    print(f"REPORT: {report_path}", flush=True)


if __name__ == "__main__":
    main()

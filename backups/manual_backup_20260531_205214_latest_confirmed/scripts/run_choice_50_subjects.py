import asyncio
import importlib.util
import json
import time
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


WORKSPACE = Path(__file__).resolve().parents[1]
EXCEL_PATH = Path.home() / "Desktop" / "企业名单.xlsx"
CHOICE_SKILL = Path.home() / ".codex" / "skills" / "mx-finance-search" / "scripts" / "get_data.py"
OUTPUT_DIR = WORKSPACE / "choice_runs"
MIN_SUBMIT_INTERVAL_SECONDS = 0.55
NEGATIVE_KEYWORDS = [
    "失信",
    "被执行",
    "限制高消费",
    "限高",
    "违约",
    "展期",
    "兑付",
    "逾期",
    "破产",
    "重整",
    "清算",
    "重大诉讼",
    "仲裁",
    "冻结",
    "查封",
    "评级下调",
    "评级关注",
    "行政处罚",
    "监管处罚",
    "债务",
    "偿债",
    "担保履约",
]


def load_choice_module():
    spec = importlib.util.spec_from_file_location("choice_get_data", CHOICE_SKILL)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def load_first_entities(limit=50):
    workbook = load_workbook(EXCEL_PATH, data_only=True, read_only=True)
    worksheet = workbook.active
    entities = OrderedDict()
    for project, borrower, guarantor in worksheet.iter_rows(min_row=2, max_col=3, values_only=True):
        for role, name in (("借款人", borrower), ("保证人", guarantor)):
            if not name:
                continue
            name = str(name).strip()
            if name in ("#N/A", "N/A"):
                continue
            item = entities.setdefault(name, {"name": name, "roles": set(), "projects": []})
            item["roles"].add(role)
            if project:
                item["projects"].append(str(project).strip())
        if len(entities) >= limit:
            break
    result = []
    for item in list(entities.values())[:limit]:
        result.append({
            "name": item["name"],
            "roles": sorted(item["roles"]),
            "projects": item["projects"][:5],
        })
    return result


def build_query(entity):
    return (
        f"请检索近三个月主体【{entity['name']}】是否存在可能严重影响还款意愿或担保履约意愿的重大负面舆情。"
        "只关注失信被执行、限制高消费、重大被执行、债务违约、展期、兑付困难、破产重整、重大诉讼仲裁、资产冻结、"
        "重大行政/监管处罚、评级下调或评级关注、担保履约能力明显恶化等事项。"
        "如果没有重大负面舆情，请明确回答未发现重大负面舆情。"
        "如有，请列出事件、发布时间、影响判断和信息来源。"
    )


def quick_has_negative(content):
    return any(keyword in content for keyword in NEGATIVE_KEYWORDS)


async def run(limit=50, test_one=False):
    choice = load_choice_module()
    entities = load_first_entities(1 if test_one else limit)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    jsonl_path = OUTPUT_DIR / f"choice_negative_news_{stamp}.jsonl"
    summary_path = OUTPUT_DIR / f"choice_negative_news_summary_{stamp}.md"

    last_submit_at = 0.0
    rows = []
    with jsonl_path.open("w", encoding="utf-8") as handle:
        for index, entity in enumerate(entities, start=1):
            now = time.monotonic()
            wait_seconds = MIN_SUBMIT_INTERVAL_SECONDS - (now - last_submit_at)
            if wait_seconds > 0:
                await asyncio.sleep(wait_seconds)
            last_submit_at = time.monotonic()

            query = build_query(entity)
            print(f"[{index}/{len(entities)}] 查询 {entity['name']}", flush=True)
            started = datetime.now().isoformat(timespec="seconds")
            result = await choice.query_financial_news(
                query=query,
                output_dir=OUTPUT_DIR / "raw",
                save_to_file=True,
            )
            content = result.get("content") or ""
            row = {
                "index": index,
                "entity": entity,
                "query": query,
                "started_at": started,
                "finished_at": datetime.now().isoformat(timespec="seconds"),
                "has_error": "error" in result,
                "error": result.get("error"),
                "output_path": result.get("output_path"),
                "rough_negative_keyword_hit": quick_has_negative(content),
                "content_preview": content[:1200],
            }
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")
            handle.flush()
            rows.append(row)

    with summary_path.open("w", encoding="utf-8") as handle:
        handle.write("# Choice数据库 50主体负面舆情试运行\n\n")
        handle.write(f"- 主体数量：{len(rows)}\n")
        handle.write(f"- 提交限速：每次提交间隔不少于 {MIN_SUBMIT_INTERVAL_SECONDS} 秒\n")
        handle.write(f"- JSONL明细：{jsonl_path}\n\n")
        handle.write("| 序号 | 主体 | 角色 | 粗略关键词命中 | 错误 | 预览 |\n")
        handle.write("|---:|---|---|---|---|---|\n")
        for row in rows:
            preview = (row["content_preview"] or "").replace("\n", " ").replace("|", "｜")[:180]
            handle.write(
                f"| {row['index']} | {row['entity']['name']} | {','.join(row['entity']['roles'])} | "
                f"{'是' if row['rough_negative_keyword_hit'] else '否'} | {row['error'] or ''} | {preview} |\n"
            )

    print(f"JSONL: {jsonl_path}")
    print(f"SUMMARY: {summary_path}")


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=50)
    parser.add_argument("--test-one", action="store_true")
    args = parser.parse_args()
    asyncio.run(run(limit=args.limit, test_one=args.test_one))
